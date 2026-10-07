"""AI Vocabulary Extraction Service using Google Gemini 2.5 Flash.
Specialized for bilingual vocabulary lists, diplomatic/technical documents, tables, and raw texts.
"""
import os
import re
import json
import base64
import logging
from typing import List, Dict, Any, Optional
import httpx
from dotenv import load_dotenv

load_dotenv(override=True)
logger = logging.getLogger(__name__)

def get_gemini_api_key() -> str:
    """Dynamically get GEMINI_API_KEY from environment or .env, reloading on demand."""
    load_dotenv(override=True)
    return (os.getenv("GEMINI_API_KEY", "") or "").strip()

def get_groq_api_key() -> str:
    load_dotenv(override=True)
    return (os.getenv("GROQ_API_KEY", "") or "").strip()


class AIVocabService:
    @staticmethod
    def extract_text_from_file(
        file_bytes: bytes,
        filename: str,
        start_page: Optional[int] = None,
        end_page: Optional[int] = None
    ) -> str:
        """Extract plain text from PDF, DOCX, TXT, or Markdown file bytes with multiple engine fallbacks.
        Supports selective page range (start_page to end_page, 1-indexed) for PDF documents.
        """
        fname_lower = filename.lower()
        if fname_lower.endswith(".pdf"):
            extracted_pages = []
            has_pdf_engine = False

            # Determine 0-based page bounds if requested
            s_idx = max(0, (start_page - 1) if (start_page and start_page > 0) else 0)

            # Engine 1: Try PyMuPDF (fitz)
            try:
                import fitz
                has_pdf_engine = True
                doc = fitz.open(stream=file_bytes, filetype="pdf")
                total_p = len(doc)
                e_idx = min(total_p, end_page if (end_page and end_page > 0) else total_p)
                if s_idx >= total_p:
                    s_idx = 0
                for i in range(s_idx, e_idx):
                    p_text = doc[i].get_text()
                    if p_text and p_text.strip():
                        extracted_pages.append(p_text)
                if extracted_pages:
                    logger.info(f"Extracted {len(extracted_pages)} pages (range {s_idx+1}-{e_idx} of {total_p}) using PyMuPDF.")
                    return "\n".join(extracted_pages).strip()
            except ImportError:
                logger.debug("PyMuPDF (fitz) is not installed, trying pypdf...")
            except Exception as e:
                logger.warning(f"PyMuPDF extraction failed: {e}")

            # Engine 2: Try pypdf (pure Python fallback)
            try:
                import pypdf
                has_pdf_engine = True
                import io
                reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                total_p = len(reader.pages)
                e_idx = min(total_p, end_page if (end_page and end_page > 0) else total_p)
                if s_idx >= total_p:
                    s_idx = 0
                for i in range(s_idx, e_idx):
                    p_text = reader.pages[i].extract_text()
                    if p_text and p_text.strip():
                        extracted_pages.append(p_text)
                if extracted_pages:
                    logger.info(f"Extracted {len(extracted_pages)} pages (range {s_idx+1}-{e_idx} of {total_p}) using pypdf.")
                    return "\n".join(extracted_pages).strip()
            except ImportError:
                logger.debug("pypdf is not installed, trying pdfminer...")
            except Exception as e:
                logger.warning(f"pypdf extraction failed: {e}")

            # Engine 3: Try pdfminer.six
            try:
                import pdfminer.high_level
                has_pdf_engine = True
                import io
                page_numbers = None
                if start_page or end_page:
                    page_numbers = list(range(s_idx, end_page if end_page else s_idx + 20))
                text = pdfminer.high_level.extract_text(io.BytesIO(file_bytes), page_numbers=page_numbers)
                if text and text.strip():
                    return text.strip()
            except ImportError:
                pass
            except Exception as e:
                logger.warning(f"pdfminer extraction failed: {e}")

            if not extracted_pages:
                if has_pdf_engine:
                    raise RuntimeError(
                        "Tệp PDF không chứa lớp văn bản (text) có thể đọc được (có thể là tệp scan dạng ảnh hoặc trang bìa không chứa chữ). "
                        "Vui lòng chọn trang khác hoặc dán trực tiếp nội dung văn bản vào ô Dán Văn Bản."
                    )
                else:
                    raise RuntimeError(
                        "Máy chủ chưa cài thư viện đọc PDF (PyMuPDF hoặc pypdf). "
                        "Vui lòng cài 'pip install pypdf PyMuPDF' trên máy chủ hoặc dán trực tiếp nội dung văn bản vào ô Dán Văn Bản."
                    )

        elif fname_lower.endswith(".docx"):
            try:
                import docx
                import io
                doc = docx.Document(io.BytesIO(file_bytes))
                lines = [p.text for p in doc.paragraphs if p.text.strip()]
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                        if row_text:
                            lines.append(row_text)
                return "\n".join(lines).strip()
            except ImportError:
                raise RuntimeError(
                    "Máy chủ chưa cài thư viện đọc Word (python-docx). "
                    "Vui lòng cài 'pip install python-docx' trên máy chủ hoặc dán trực tiếp nội dung văn bản vào ô Dán Văn Bản."
                )
        
        else:
            # Assume text/markdown/csv
            for enc in ["utf-8", "utf-16", "cp1252", "latin-1"]:
                try:
                    return file_bytes.decode(enc).strip()
                except UnicodeDecodeError:
                    continue
            return file_bytes.decode("utf-8", errors="replace").strip()

    @classmethod
    async def extract_vocabulary(
        cls,
        text: str = "",
        source_lang: str = "en",
        target_lang: str = "vi",
        mode: str = "auto",
        vocab_level: str = "intermediate",
        images: Optional[List[bytes]] = None
    ) -> List[Dict[str, Any]]:
        """Call Gemini to extract, normalize, and enrich vocabulary pairs from ANY document text or images/scanned PDFs."""
        if (not text or not text.strip()) and not images:
            return []

        # Enforce maximum 5,000 words limit for text
        cleaned_input = ""
        if text and text.strip():
            words = text.strip().split()
            if len(words) > 5000:
                cleaned_input = " ".join(words[:5000])
                logger.info(f"Input text exceeded 5,000 words ({len(words)} words). Truncated to first 5,000 words.")
            else:
                cleaned_input = text.strip()

        # Vocabulary Level Filter Specs
        LEVEL_SPECS = {
            "all": {
                "name": "Mọi cấp độ (A1 – C2)",
                "instruction": "Trích xuất các từ vựng và cụm từ hữu ích từ bài đọc, từ cơ bản đến nâng cao.",
                "exclude": "Chỉ bỏ qua các mạo từ đơn lẻ (a, an, the) hoặc đại từ nhân xưng vô nghĩa (I, you, he, she)."
            },
            "intermediate": {
                "name": "Trung cấp trở lên (B1 – C2 / IELTS 5.0 - 6.5)",
                "instruction": "Chỉ chọn lọc từ vựng từ cấp độ B1 trở lên, collocations, cụm động từ (phrasal verbs), thành ngữ (idioms) và từ vựng mang tính học tập.",
                "exclude": (
                    "TUYỆT ĐỐI KHÔNG LẤY các từ vựng sơ cấp A1 - A2 quá dễ và hiển nhiên. "
                    "Ví dụ BẮT BUỘC BỎ: say, note, tell, ask, speak, talk, see, look, go, come, make, do, have, get, "
                    "give, take, good, bad, happy, sad, big, small, new, old, very, really, because, but, and, so, or..."
                )
            },
            "advanced": {
                "name": "Nâng cao & Học thuật (B2 – C2 / IELTS 6.5 - 7.5+)",
                "instruction": "Chỉ chọn lọc các từ vựng học thuật (academic), thuật ngữ báo chí/khoa học, collocations đắt giá, thành ngữ và từ vựng band cao (ví dụ: underscore, highlight, stress, corroborate, exacerbate, unprecedented...).",
                "exclude": (
                    "TUYỆT ĐỐI LOẠI BỎ toàn bộ từ vựng dưới cấp độ B2 (B1, A2, A1). "
                    "Tuyệt đối không lấy các từ thông dụng như: say, note, tell, important, problem, result, effect, different, difficult, easy..."
                )
            },
            "expert": {
                "name": "Chuyên sâu / C1 – C2 (IELTS 8.0+)",
                "instruction": "Chỉ chọn các từ vựng C1 – C2 cao cấp, thuật ngữ chuyên ngành tinh tế, từ vựng văn phong học thuật xuất sắc và cấu trúc hiếm gặp.",
                "exclude": "TUYỆT ĐỐI LOẠI BỎ toàn bộ từ vựng phổ thông các cấp độ A1, A2, B1 và B2."
            }
        }
        lvl = LEVEL_SPECS.get(vocab_level, LEVEL_SPECS["intermediate"])

        # Language-specific descriptors for 6 supported languages
        LANG_SPECS = {
            "en": {
                "name": "Tiếng Anh",
                "word_rule": "Từ hoặc cụm từ TIẾNG ANH chuẩn (nguyên thể hoặc collocations, phrasal verbs, idioms)",
                "phonetic_rule": "Phiên âm quốc tế IPA chuẩn của từ/cụm từ (ví dụ: '/ˈpræk.tɪs/')",
                "pos_rule": "Loại từ (noun, verb, adjective, adverb, phrasal verb, idiom, phrase...)",
                "sample_word": "artificial intelligence",
                "sample_meaning": "trí tuệ nhân tạo",
                "sample_phonetic": "/ˌɑː.tɪ.fɪʃ.əl ɪnˈtel.ɪ.dʒəns/",
                "sample_pos": "noun phrase"
            },
            "ja": {
                "name": "Tiếng Nhật",
                "word_rule": "Từ hoặc cụm từ TIẾNG NHẬT (chữ Hán Kanji hoặc Hiragana/Katakana chuẩn)",
                "phonetic_rule": "Cách đọc Hiragana / Furigana và Romaji (ví dụ: 'かんじ (kanji)')",
                "pos_rule": "Loại từ (Danh từ, Động từ nhóm 1/2/3, Tính từ đuôi -i, Tính từ đuôi -na, Phó từ, Cụm từ...)",
                "sample_word": "人工知能",
                "sample_meaning": "trí tuệ nhân tạo",
                "sample_phonetic": "じんこうちのう (jinkō chinō)",
                "sample_pos": "Danh từ"
            },
            "zh": {
                "name": "Tiếng Trung",
                "word_rule": "Từ hoặc cụm từ TIẾNG TRUNG (chữ Hán giản thể hoặc phồn thể)",
                "phonetic_rule": "Bính âm Pinyin có đầy đủ dấu thanh điệu (ví dụ: 'rén gōng zhì néng')",
                "pos_rule": "Loại từ (Danh từ, Động từ, Tính từ, Lượng từ, Phó từ, Thành ngữ...)",
                "sample_word": "人工智能",
                "sample_meaning": "trí tuệ nhân tạo",
                "sample_phonetic": "rén gōng zhì néng",
                "sample_pos": "Danh từ"
            },
            "ko": {
                "name": "Tiếng Hàn",
                "word_rule": "Từ hoặc cụm từ TIẾNG HÀN (chữ Hangul chuẩn)",
                "phonetic_rule": "Phiên âm Romaja / Romanization chuẩn (ví dụ: 'in-gong-ji-neung')",
                "pos_rule": "Loại từ (Danh từ, Động từ, Tính từ, Phó từ, Cụm từ...)",
                "sample_word": "인공지능",
                "sample_meaning": "trí tuệ nhân tạo",
                "sample_phonetic": "in-gong-ji-neung",
                "sample_pos": "Danh từ"
            },
            "fr": {
                "name": "Tiếng Pháp",
                "word_rule": "Từ hoặc cụm từ TIẾNG PHÁP chuẩn (kèm mạo từ le/la nếu là danh từ)",
                "phonetic_rule": "Phiên âm quốc tế IPA chuẩn của từ (ví dụ: '/ɛ̃.tɛ.li.ʒɑ̃s/')",
                "pos_rule": "Loại từ (nom masculin, nom féminin, verbe, adjectif, adverbe, locution...)",
                "sample_word": "intelligence artificielle",
                "sample_meaning": "trí tuệ nhân tạo",
                "sample_phonetic": "/ɛ̃.tɛ.li.ʒɑ̃s aʁ.ti.fi.sjɛl/",
                "sample_pos": "nom féminin"
            },
            "de": {
                "name": "Tiếng Đức",
                "word_rule": "Từ hoặc cụm từ TIẾNG ĐỨC chuẩn (viết hoa danh từ kèm mạo từ der/die/das)",
                "phonetic_rule": "Phiên âm quốc tế IPA chuẩn của từ (ví dụ: '/ˈhaʊs/')",
                "pos_rule": "Loại từ (Nomen maskulin/feminin/neutral, Verb, Adjektiv, Adverb, Redewendung...)",
                "sample_word": "künstliche Intelligenz",
                "sample_meaning": "trí tuệ nhân tạo",
                "sample_phonetic": "/ˈkʏnstlɪçə ɪntɛliˈɡɛnts/",
                "sample_pos": "Nomen (feminin)"
            }
        }
        spec = LANG_SPECS.get(source_lang, LANG_SPECS["en"])
        lang_name = spec["name"]

        prompt = f"""Bạn là một chuyên gia ngôn ngữ học và giảng dạy {lang_name} hàng đầu.
Nhiệm vụ của bạn là đọc và phân tích KỸ LƯỠNG tài liệu/văn bản {lang_name} dưới đây (kể cả tệp scan, báo cáo, đề thi, danh sách từ vựng song ngữ, hoặc bài đọc chuyên ngành).

Hãy tự động nhận diện và trích xuất TOÀN BỘ các từ vựng quan trọng, cụm từ đắt giá, cấu trúc câu hay, và thuật ngữ hữu ích cho người học.

QUY CHUẨN CẤP ĐỘ BẮT BUỘC: {lvl['name']}
- Mục tiêu tuyển chọn: {lvl['instruction']}
- {lvl['exclude']}

Yêu cầu bóc tách chi tiết:
1. Áp dụng linh hoạt cho MỌI thể loại văn bản:
   - Nếu là đoạn văn bản tự do: Chọn lọc các từ vựng cốt lõi, từ vựng nâng cao, thành ngữ và cụm từ xuất hiện trong bài đạt chuẩn {lvl['name']}.
   - Nếu là danh sách từ vựng/bảng song ngữ: Tách chính xác từng cặp từ {lang_name} và nghĩa tiếng Việt tương ứng.
2. Tách rõ ràng từng mục:
   - "word": {spec['word_rule']}.
   - "meaning": Nghĩa TIẾNG VIỆT chính xác, tự nhiên, sát ngữ cảnh của bài.
   - "phonetic": {spec['phonetic_rule']}.
   - "part_of_speech": {spec['pos_rule']}.
   - "context_sentence": Câu ví dụ trích trực tiếp từ văn bản đầu vào (hoặc câu ví dụ tự nhiên minh họa cách dùng từ này).
3. Định dạng trả về BẮT BUỘC là JSON Array thuần túy chứa danh sách các object, ví dụ:
[
  {{
    "word": "{spec['sample_word']}",
    "meaning": "{spec['sample_meaning']}",
    "phonetic": "{spec['sample_phonetic']}",
    "part_of_speech": "{spec['sample_pos']}",
    "context_sentence": "Ví dụ minh họa ngữ cảnh câu chứa từ này."
  }}
]
"""
        if cleaned_input:
            prompt += f"""
VĂN BẢN ĐẦU VÀO ({lang_name}):
\"\"\"
{cleaned_input}
\"\"\"
"""
        api_key = get_gemini_api_key()
        if not api_key:
            raise RuntimeError("GEMINI_API_KEY chưa được cấu hình trên máy chủ.")

        models_to_try = ["gemini-2.5-flash", "gemini-1.5-flash"]
        last_error = None

        # Build multimodal contents parts
        contents_parts = []
        if images:
            for img in images:
                contents_parts.append({
                    "inline_data": {
                        "mime_type": "image/jpeg",
                        "data": base64.b64encode(img).decode("utf-8")
                    }
                })
            contents_parts.append({
                "text": f"TÀI LIỆU HÌNH ẢNH / TRANG SCAN ĐÍNH KÈM:\nHãy đọc kỹ và nhận diện toàn bộ văn bản xuất hiện trong các bức ảnh này (OCR).\nSau đó thực hiện yêu cầu bóc tách từ vựng dưới đây:\n\n{prompt}"
            })
        else:
            contents_parts.append({"text": prompt})

        payload = {
            "contents": [{"parts": contents_parts}],
            "generationConfig": {
                "responseMimeType": "application/json",
                "temperature": 0.2
            }
        }

        async with httpx.AsyncClient(timeout=300.0) as client:
            for model_name in models_to_try:
                url = f"https://generativelanguage.googleapis.com/v1beta/models/{model_name}:generateContent?key={api_key}"
                try:
                    response = await client.post(url, json=payload)
                    if response.status_code == 200:
                        res_json = response.json()
                        candidates = res_json.get("candidates", [])
                        if candidates and "content" in candidates[0]:
                            parts = candidates[0]["content"].get("parts", [])
                            if parts and "text" in parts[0]:
                                raw_text = parts[0]["text"].strip()
                                # Clean markdown code fences if present
                                raw_text = re.sub(r"^```(?:json)?\s*", "", raw_text, flags=re.MULTILINE)
                                raw_text = re.sub(r"\s*```$", "", raw_text, flags=re.MULTILINE)
                                parsed_list = json.loads(raw_text)
                                if isinstance(parsed_list, list):
                                    # Normalize each item
                                    results = []
                                    seen_words = set()
                                    for item in parsed_list:
                                        if not isinstance(item, dict):
                                            continue
                                        w = str(item.get("word", "")).strip()
                                        m = str(item.get("meaning", "")).strip()
                                        if not w:
                                            continue
                                        if w.lower() in seen_words:
                                            continue
                                        seen_words.add(w.lower())
                                        results.append({
                                            "word": w,
                                            "meaning": m or w,
                                            "phonetic": str(item.get("phonetic", "")).strip() or None,
                                            "part_of_speech": str(item.get("part_of_speech", "phrase")).strip(),
                                            "context_sentence": str(item.get("context_sentence", "")).strip()
                                        })
                                    return results
                    else:
                        last_error = f"Gemini API ({model_name}) error HTTP {response.status_code}: {response.text}"
                        logger.warning(last_error)
                except Exception as e:
                    last_error = str(e)
                    logger.warning(f"Error calling {model_name}: {e}")

        raise RuntimeError(f"Không thể kết nối hoặc bóc tách từ AI: {last_error}")
